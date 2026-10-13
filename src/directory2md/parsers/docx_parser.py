"""Microsoft Word (.docx) document parser using python-docx.

Extracts text paragraphs and formatted table cells from Word documents.
"""

from __future__ import annotations

from pathlib import Path
from typing import Tuple

from directory2md.models import DocumentMetadata
from directory2md.parsers.base import BaseParser


class DocxParser(BaseParser):
    """Parser implementation for Microsoft Word (.docx) files."""

    def parse(self, filepath: Path) -> Tuple[DocumentMetadata, str]:
        """Extract paragraphs and table text from a Word document formatted as Markdown.

        Args:
            filepath: Path to the target .docx file.

        Returns:
            Tuple of (:class:`DocumentMetadata`, extracted_full_text).
        """
        metadata = self.build_initial_metadata(filepath)
        text_parts = []

        try:
            from docx import Document

            doc = Document(str(filepath))
            for para in doc.paragraphs:
                clean_p = para.text.strip()
                if clean_p:
                    # Check if paragraph has heading style
                    if para.style and para.style.name and para.style.name.lower().startswith("heading 1"):
                        text_parts.append(f"# {clean_p}")
                    elif para.style and para.style.name and para.style.name.lower().startswith("heading 2"):
                        text_parts.append(f"## {clean_p}")
                    elif para.style and para.style.name and para.style.name.lower().startswith("heading"):
                        text_parts.append(f"### {clean_p}")
                    else:
                        text_parts.append(clean_p)

            for table in doc.tables:
                table_lines = []
                is_first = True
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
                    if any(cells):
                        table_lines.append(f"| {' | '.join(cells)} |")
                        if is_first:
                            table_lines.append(f"| {' | '.join([':---'] * len(cells))} |")
                            is_first = False
                if table_lines:
                    text_parts.append("\n" + "\n".join(table_lines) + "\n")
        except Exception as e:
            text_parts.append(f"[Error reading DOCX: {e}]")

        full_text = "\n\n".join(text_parts).strip()
        words = full_text.split()
        metadata.word_count = len(words)
        metadata.text_snippet = " ".join(words[:60]) + ("..." if len(words) > 60 else "")
        metadata.extracted_text = full_text

        return metadata, full_text

